from __future__ import annotations

import csv
import math
from pathlib import Path
from statistics import fmean, pstdev
from typing import Any

from yzwcloud.plot_studio_utils import (
    _first_present,
    _normalize_column_name,
    _parse_float,
    _round_number,
)


def inspect_table(path: Path, *, max_rows: int = 2000) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        header, rows, truncated = _read_xlsx_rows(path, max_rows=max_rows)
        file_format = "xlsx"
    else:
        header, rows, truncated = _read_delimited_rows(path, max_rows=max_rows)
        file_format = "delimited"

    columns = _clean_header(header)
    column_stats = _collect_column_stats(columns, rows)
    numeric_columns = [name for name, stats in column_stats.items() if stats["kind"] == "numeric"]
    categorical_columns = [name for name, stats in column_stats.items() if stats["kind"] != "numeric"]
    signals = _detect_table_signals(columns, rows)

    return {
        "path": str(path),
        "filename": path.name,
        "format": file_format,
        "scanned_rows": len(rows),
        "truncated": truncated,
        "column_count": len(columns),
        "columns": columns,
        "numeric_columns": numeric_columns,
        "categorical_columns": categorical_columns,
        "column_summaries": column_stats,
        "signals": signals,
    }


def _read_delimited_rows(path: Path, *, max_rows: int) -> tuple[list[str], list[list[str]], bool]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        sample = file.read(4096)
        file.seek(0)
        delimiter = _infer_delimiter(path, sample)
        reader = csv.reader(file, delimiter=delimiter)
        header = next(reader, [])
        rows = []
        truncated = False
        for index, row in enumerate(reader):
            if index >= max_rows:
                truncated = True
                break
            rows.append([str(value) for value in row])
    return [str(value) for value in header], rows, truncated


def _load_table_records(path: Path, *, max_rows: int) -> tuple[list[str], list[dict[str, str]]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        header, rows, _ = _read_xlsx_rows(path, max_rows=max_rows)
    else:
        header, rows, _ = _read_delimited_rows(path, max_rows=max_rows)
    columns = _clean_header(header)
    records = []
    for row in rows:
        records.append({column: row[index] if index < len(row) else "" for index, column in enumerate(columns)})
    return columns, records


def _read_xlsx_rows(path: Path, *, max_rows: int) -> tuple[list[str], list[list[str]], bool]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    header = ["" if value is None else str(value) for value in next(iterator, [])]
    rows: list[list[str]] = []
    truncated = False
    for index, row in enumerate(iterator):
        if index >= max_rows:
            truncated = True
            break
        rows.append(["" if value is None else str(value) for value in row])
    workbook.close()
    return header, rows, truncated


def _infer_delimiter(path: Path, sample: str) -> str:
    if path.suffix.lower() == ".tsv":
        return "\t"
    if path.suffix.lower() == ".csv":
        return ","
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        return ","
    return str(dialect.delimiter)


def _clean_header(header: list[str]) -> list[str]:
    used: dict[str, int] = {}
    columns = []
    for index, raw_name in enumerate(header):
        name = str(raw_name or "").strip() or f"column_{index + 1}"
        count = used.get(name, 0) + 1
        used[name] = count
        if count > 1:
            name = f"{name}_{count}"
        columns.append(name)
    return columns


def _collect_column_stats(columns: list[str], rows: list[list[str]]) -> dict[str, dict[str, Any]]:
    raw_stats: dict[str, dict[str, Any]] = {
        name: {"missing": 0, "non_missing": 0, "numeric_values": [], "distinct": set()}
        for name in columns
    }
    for row in rows:
        for index, column in enumerate(columns):
            value = row[index] if index < len(row) else ""
            text = str(value).strip()
            stats = raw_stats[column]
            if text == "":
                stats["missing"] += 1
                continue
            stats["non_missing"] += 1
            if len(stats["distinct"]) < 25:
                stats["distinct"].add(text)
            parsed = _parse_float(text)
            if parsed is not None and math.isfinite(parsed):
                stats["numeric_values"].append(parsed)

    summaries: dict[str, dict[str, Any]] = {}
    for column, stats in raw_stats.items():
        values = stats["numeric_values"]
        non_missing = int(stats["non_missing"])
        numeric_ratio = len(values) / non_missing if non_missing else 0.0
        if values and numeric_ratio >= 0.8:
            summaries[column] = {
                "kind": "numeric",
                "count": len(values),
                "missing": int(stats["missing"]),
                "min": _round_number(min(values)),
                "max": _round_number(max(values)),
                "mean": _round_number(fmean(values)),
                "std": _round_number(pstdev(values)) if len(values) > 1 else 0,
            }
        else:
            summaries[column] = {
                "kind": "categorical",
                "count": non_missing,
                "missing": int(stats["missing"]),
                "unique_preview": sorted(str(item) for item in stats["distinct"]),
            }
    return summaries


def _detect_table_signals(columns: list[str], rows: list[list[str]]) -> dict[str, Any]:
    normalized = {_normalize_column_name(column): index for index, column in enumerate(columns)}
    signals: dict[str, Any] = {}

    log2fc_index = _first_present(normalized, "log2fc", "log2_fold_change")
    p_index = _first_present(normalized, "p_value", "pvalue", "padj")
    if log2fc_index is not None and p_index is not None:
        significant = 0
        up = 0
        down = 0
        for row in rows:
            if log2fc_index >= len(row) or p_index >= len(row):
                continue
            log2fc = _parse_float(row[log2fc_index])
            p_value = _parse_float(row[p_index])
            if log2fc is None or p_value is None:
                continue
            if abs(log2fc) >= 1.0 and p_value <= 0.05:
                significant += 1
                if log2fc > 0:
                    up += 1
                elif log2fc < 0:
                    down += 1
        signals["differential_default_threshold"] = {
            "abs_log2fc": 1.0,
            "p_value": 0.05,
            "significant": significant,
            "up": up,
            "down": down,
        }

    numeric_columns = _numeric_columns_from_rows(columns, rows)
    identifier_columns = [
        column
        for column in columns
        if _normalize_column_name(column)
        in {
            "gene",
            "gene_id",
            "gene_name",
            "gene_short_name",
            "symbol",
            "feature",
            "feature_id",
            "ensembl_id",
            "transcript_id",
        }
    ]
    excluded_numeric_columns = [column for column in numeric_columns if _is_matrix_numeric_metadata(column)]
    value_columns = [column for column in numeric_columns if column not in excluded_numeric_columns]
    if identifier_columns and len(value_columns) >= 6:
        signals["matrix_profile"] = {
            "kind": "expression_like",
            "identifier_columns": identifier_columns[:5],
            "value_columns": value_columns,
            "excluded_numeric_columns": excluded_numeric_columns,
            "numeric_value_count": len(value_columns),
        }
    return signals


def _numeric_columns_from_rows(columns: list[str], rows: list[list[str]]) -> list[str]:
    numeric_columns = []
    for index, column in enumerate(columns):
        numeric_count = 0
        non_missing = 0
        for row in rows:
            if index >= len(row):
                continue
            text = str(row[index]).strip()
            if not text:
                continue
            non_missing += 1
            parsed = _parse_float(text)
            if parsed is not None and math.isfinite(parsed):
                numeric_count += 1
        if non_missing and numeric_count / non_missing >= 0.8:
            numeric_columns.append(column)
    return numeric_columns


def _is_matrix_numeric_metadata(column: str) -> bool:
    normalized = _normalize_column_name(column)
    if normalized in {
        "length",
        "gene_length",
        "transcript_length",
        "tx_length",
        "width",
        "start",
        "end",
        "chrom_start",
        "chrom_end",
        "tx_start",
        "tx_end",
        "cds_start",
        "cds_end",
    }:
        return True
    return normalized.endswith("_start") or normalized.endswith("_end") or "length" in normalized


