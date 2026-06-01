from __future__ import annotations

import csv
import gzip
import json
import os
import shutil
import tarfile
import time
import urllib.error
import urllib.request
import zipfile
from collections.abc import Iterable
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from yzwcloud.config import PROJECT_ROOT
from yzwcloud.color_palette import condition_color_map
from yzwcloud.data_intake_registry import (
    DEFAULT_MAX_ITERATIONS,
    attempt_stop_reason,
    build_processing_plan,
    direct_capabilities_for_data_type,
    execute_processing_strategy,
    plan_stop_message,
)
from yzwcloud.models import DataObject
from yzwcloud.prompts import DATA_INTAKE_SYSTEM_PROMPT


GENE_COLUMNS = ["gene_short_name", "gene_id", "biotype", "strand", "locus", "Length"]
REQUIRED_MATRIX_GENE_COLUMNS = 6
DEFAULT_MODEL = "deepseek-v4-flash"


def run_data_intake_agent(
    source_path: Path,
    output_dir: Path,
    metadata_path: Path | None = None,
    params: dict[str, Any] | None = None,
    progress_callback: Callable[[str, str, str], None] | None = None,
) -> DataObject:
    params = params or {}
    output_dir.mkdir(parents=True, exist_ok=True)
    checkpoint_path = output_dir / "data_intake_checkpoint.json"
    checkpoint = {
        "agent": "data_intake_agent",
        "status": "running",
        "source_file": str(source_path),
        "metadata_file": str(metadata_path) if metadata_path else "",
        "iterations": [],
        "attempts": [],
        "created_at": datetime.now().isoformat(),
        "updated_at": datetime.now().isoformat(),
    }
    _write_checkpoint(checkpoint_path, checkpoint)

    try:
        max_iterations = max(1, int(params.get("max_iterations", DEFAULT_MAX_ITERATIONS)))
        _emit_progress(progress_callback, "prepare_input", "running", "正在展开输入")
        prepared = _prepare_source_input(source_path, output_dir)
        checkpoint["prepared_input"] = prepared
        _record_step(checkpoint_path, checkpoint, "prepare_input", "completed", prepared)
        _emit_progress(progress_callback, "prepare_input", "completed", "已展开输入")

        _emit_progress(progress_callback, "inspect_file", "running", "正在读取数据")
        prepared_source_path = Path(prepared["resolved_source_path"])
        inspection = _inspect_source(prepared_source_path, metadata_path)
        inspection["original_source_file"] = str(source_path)
        inspection["resolved_source_file"] = str(prepared_source_path)
        inspection["input_preprocess"] = prepared
        _record_step(checkpoint_path, checkpoint, "inspect_file", "completed", inspection)
        _emit_progress(progress_callback, "inspect_file", "completed", "已读取数据")

        _emit_progress(progress_callback, "classify_data", "running", "正在识别类型")
        llm_plan = _classify_with_deepseek(inspection, use_llm=bool(params.get("use_llm", True)))
        _record_step(checkpoint_path, checkpoint, "classify_data", "completed", llm_plan)
        _emit_progress(progress_callback, "classify_data", "completed", "已识别类型")
        plan = build_processing_plan(inspection, llm_plan, metadata_path, max_iterations)
        checkpoint["plan"] = plan
        _record_step(checkpoint_path, checkpoint, "plan_processing", "completed", plan)

        if not plan["strategies"]:
            raise ValueError(plan_stop_message(plan))

        standard_result = None
        validation = None
        success_strategy = None
        last_error = ""
        for attempt_index, strategy in enumerate(plan["strategies"][:max_iterations], start=1):
            attempt_label = f"第 {attempt_index} 轮：{strategy['label']}"
            _emit_progress(progress_callback, "standardize_data", "running", attempt_label)
            attempt = {
                "attempt": attempt_index,
                "strategy_id": strategy["id"],
                "label": strategy["label"],
                "status": "running",
                "started_at": datetime.now().isoformat(),
            }
            try:
                standard_result = execute_processing_strategy(
                    source_path=prepared_source_path,
                    metadata_path=metadata_path,
                    output_dir=output_dir,
                    inspection=inspection,
                    llm_plan=llm_plan,
                    strategy=strategy,
                    standardizers={
                        "expression_matrix": _standardize_expression_like_table,
                        "metabolomics_matrix": _standardize_metabolomics_table,
                    },
                )
                _emit_progress(progress_callback, "standardize_data", "completed", f"{attempt_label} 已规整")

                _emit_progress(progress_callback, "validate_output", "running", f"{attempt_label} 正在验证")
                validation = _validate_expression_matrix(
                    matrix_path=Path(standard_result["matrix_file"]),
                    metadata_path=Path(standard_result["sample_metadata_file"]),
                )
                attempt["validation"] = validation
                if validation["valid"]:
                    attempt["status"] = "completed"
                    attempt["completed_at"] = datetime.now().isoformat()
                    _record_attempt(checkpoint_path, checkpoint, attempt)
                    success_strategy = strategy
                    _record_step(checkpoint_path, checkpoint, "validate_output", "completed", validation)
                    _emit_progress(progress_callback, "validate_output", "completed", f"{attempt_label} 已验证")
                    break

                last_error = "; ".join(validation["errors"]) or "Data intake validation failed"
                attempt["status"] = "failed"
                attempt["error"] = last_error
                attempt["completed_at"] = datetime.now().isoformat()
                _record_attempt(checkpoint_path, checkpoint, attempt)
            except Exception as exc:
                last_error = str(exc)
                attempt["status"] = "failed"
                attempt["error"] = last_error
                attempt["completed_at"] = datetime.now().isoformat()
                _record_attempt(checkpoint_path, checkpoint, attempt)

        if success_strategy is None or standard_result is None or validation is None:
            raise ValueError(last_error or plan_stop_message(plan))
    except Exception:
        checkpoint["status"] = "failed"
        _write_checkpoint(checkpoint_path, checkpoint)
        _emit_progress(progress_callback, "failed", "failed", "处理失败，等待重试")
        raise

    standardized_data_type = str(success_strategy.get("data_type", "expression_matrix"))
    validated_capabilities = _capabilities_for_validation(validation)
    if standardized_data_type == "metabolomics_matrix" and validation["valid"]:
        validated_capabilities.append("metabolomics_normalization")
        validated_capabilities.append("metabolomics_differential")
    capabilities = direct_capabilities_for_data_type(
        data_type=standardized_data_type,
        capabilities=validated_capabilities,
    )
    next_analyses = _next_analyses_for_capabilities(capabilities)
    metadata = {
        "data_type": standardized_data_type,
        "source_file": str(source_path),
        "resolved_source_file": str(prepared_source_path),
        "sample_metadata_source": str(metadata_path) if metadata_path else "",
        "gene_count": validation["gene_count"],
        "metabolite_count": validation["gene_count"] if standardized_data_type == "metabolomics_matrix" else 0,
        "sample_count": validation["sample_count"],
        "sample_groups": validation["sample_groups"],
        "conditions": validation["conditions"],
        "condition_options": sorted(validation["conditions"]),
        "condition_colors": condition_color_map(validation["conditions"]),
        "matrix_file": standard_result["matrix_file"],
        "sample_metadata_file": standard_result["sample_metadata_file"],
        "gene_annotation_file": standard_result.get("gene_annotation_file", ""),
        "capabilities": capabilities,
        "next_analyses": next_analyses,
        "checkpoint_file": str(checkpoint_path),
        "agent": {
            "name": "data_intake_agent",
            "model": llm_plan.get("model", DEFAULT_MODEL),
            "llm_status": llm_plan.get("llm_status", "unknown"),
            "classification": llm_plan.get("data_type", "expression_matrix"),
            "warnings": validation["warnings"] + llm_plan.get("warnings", []),
            "max_iterations": max_iterations,
            "strategy_history": checkpoint.get("attempts", []),
            "selected_strategy": success_strategy,
        },
        "input_preprocess": prepared,
        "standardization": standard_result["standardization"],
        "quality": {
            "duplicate_sample_columns": inspection.get("duplicate_sample_columns", {}),
            "numeric_sample_column_ratio": inspection.get("numeric_sample_column_ratio", 0),
        },
    }
    output = DataObject(type="expression_matrix", data=standard_result["matrix_file"], meta=metadata)
    (output_dir / "data_intake_output.json").write_text(
        output.model_dump_json(indent=2),
        encoding="utf-8",
    )
    checkpoint["status"] = "completed"
    checkpoint["output"] = metadata
    _write_checkpoint(checkpoint_path, checkpoint)
    _emit_progress(progress_callback, "completed", "completed", "数据已可用于分析")
    return output


def _prepare_source_input(source_path: Path, output_dir: Path) -> dict[str, Any]:
    lower_name = source_path.name.lower()
    if lower_name.endswith((".csv", ".tsv", ".txt", ".xlsx", ".xlsm")):
        return {
            "kind": "direct_file",
            "original_source_path": str(source_path),
            "resolved_source_path": str(source_path),
        }

    unpack_dir = output_dir / "unpacked_input"
    if unpack_dir.exists():
        shutil.rmtree(unpack_dir)
    unpack_dir.mkdir(parents=True, exist_ok=True)

    if lower_name.endswith(".zip"):
        with zipfile.ZipFile(source_path) as archive:
            archive.extractall(unpack_dir)
    elif lower_name.endswith((".tar", ".tar.gz", ".tgz")):
        with tarfile.open(source_path, mode="r:*") as archive:
            archive.extractall(unpack_dir)
    elif lower_name.endswith(".gz"):
        target_name = Path(lower_name[:-3]).name or "decompressed_input"
        target_path = unpack_dir / target_name
        with gzip.open(source_path, "rb") as gz_file, target_path.open("wb") as output_file:
            shutil.copyfileobj(gz_file, output_file)
    else:
        raise ValueError(f"Unsupported uploaded file type: {source_path.name}")

    entries = sorted(path for path in unpack_dir.rglob("*") if path.is_file())
    candidate = _select_prepared_candidate(entries)
    if candidate is None:
        if _looks_like_10x_directory(entries):
            raise ValueError("压缩包里更像 10x/单细胞矩阵，当前还没有接入单细胞处理器。")
        raise ValueError("压缩包已解压，但里面没有找到当前支持的 .csv/.xlsx/.xlsm 数据文件。")

    return {
        "kind": "archive_extract",
        "original_source_path": str(source_path),
        "resolved_source_path": str(candidate),
        "unpack_dir": str(unpack_dir),
        "archive_members": [str(path.relative_to(unpack_dir)) for path in entries[:80]],
        "selected_member": str(candidate.relative_to(unpack_dir)),
    }


def _select_prepared_candidate(entries: list[Path]) -> Path | None:
    supported = [path for path in entries if path.suffix.lower() in {".csv", ".tsv", ".txt", ".xlsx", ".xlsm"}]
    if not supported:
        return None
    ranked = sorted(supported, key=_candidate_priority)
    return ranked[0]


def _candidate_priority(path: Path) -> tuple[int, int, str]:
    name = path.name.lower()
    preferred_tokens = ("maf", "metabolite", "metabolomics", "matrix", "expression", "count", "counts", "mrna", "gene")
    score = 0 if any(token in name for token in preferred_tokens) else 1
    suffix_score = 0 if path.suffix.lower() == ".csv" else 1
    return (score, suffix_score, name)


def _looks_like_10x_directory(entries: list[Path]) -> bool:
    names = {path.name.lower() for path in entries}
    return "matrix.mtx" in names and any(name.startswith("barcodes") for name in names) and any(
        name.startswith("features") or name.startswith("genes") for name in names
    )


def build_failure_report(checkpoint_path: Path, error_message: str) -> dict[str, Any]:
    if not checkpoint_path.exists():
        return {
            "title": "数据处理失败",
            "summary": "Agent 未生成检查快照，暂时无法给出更具体的结构化原因。",
            "reasons": [error_message],
            "findings": [],
            "suggestions": ["重新上传文件后再试一次。"],
        }

    checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
    iterations = checkpoint.get("iterations", [])
    attempts = checkpoint.get("attempts", [])
    plan = checkpoint.get("plan", {})
    steps = {item.get("step"): item.get("payload", {}) for item in iterations}
    inspection = steps.get("inspect_file", {})
    validation = steps.get("validate_output", {})
    findings = _failure_findings(inspection, validation)
    reasons = _failure_reasons(error_message, inspection, validation)
    suggestions = _failure_suggestions(error_message, inspection, validation)
    if attempts:
        findings.append(f"Agent 共尝试了 {len(attempts)} 种处理策略。")
        reasons.append(attempt_stop_reason(attempts, plan))
    elif plan.get("unsupported_reason"):
        reasons.append(str(plan["unsupported_reason"]))

    return {
        "title": "数据处理失败",
        "summary": _failure_summary(inspection, validation),
        "reasons": reasons,
        "findings": findings,
        "suggestions": suggestions,
        "attempts": attempts,
        "inspection": {
            "file_type": inspection.get("file_type", ""),
            "row_count": inspection.get("row_count", 0),
            "column_count": inspection.get("column_count", 0),
            "header_preview": inspection.get("header", [])[:12],
            "likely_gene_columns": inspection.get("likely_gene_columns", []),
            "numeric_column_ratio": inspection.get("numeric_column_ratio", 0),
            "numeric_sample_column_ratio": inspection.get("numeric_sample_column_ratio", 0),
            "metadata_rows": inspection.get("metadata_rows", 0),
            "duplicate_sample_columns": inspection.get("duplicate_sample_columns", {}),
        },
        "raw_error": error_message,
    }


def _inspect_source(source_path: Path, metadata_path: Path | None) -> dict[str, Any]:
    if source_path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _inspect_workbook(source_path, metadata_path)
    return _inspect_csv(source_path, metadata_path)


def _inspect_csv(source_path: Path, metadata_path: Path | None) -> dict[str, Any]:
    delimiter = _infer_delimiter(source_path)
    with source_path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file, delimiter=delimiter)
        header = next(reader)
        preview = [row for _, row in zip(range(8), reader)]

    duplicate_columns = _duplicates(header)
    numeric_ratios = _numeric_ratios(preview, len(header))
    raw_sample_metadata = _read_sample_metadata(metadata_path) if metadata_path and metadata_path.exists() else []
    sample_metadata = _dedupe_metadata_by_sample(raw_sample_metadata)
    sample_names = [row["sample"] for row in sample_metadata]
    sample_column_indices = _sample_column_indices(header, sample_names)

    return {
        "file_type": "csv",
        "delimiter": delimiter,
        "path": str(source_path),
        "column_count": len(header),
        "row_count": _count_csv_rows(source_path),
        "header": header[:160],
        "preview": preview[:3],
        "duplicate_columns": duplicate_columns,
        "duplicate_sample_columns": {
            name: indices for name, indices in sample_column_indices.items() if len(indices) > 1
        },
        "metadata_rows": len(sample_metadata),
        "metadata_conditions": _count_values(row.get("condition", "") for row in sample_metadata),
        "numeric_column_ratio": _safe_ratio(sum(ratio > 0.8 for ratio in numeric_ratios), len(numeric_ratios)),
        "numeric_sample_column_ratio": _numeric_sample_ratio(preview, sample_column_indices),
        "likely_gene_columns": [name for name in GENE_COLUMNS if name in header],
        "likely_metabolomics_columns": _likely_metabolomics_columns(header),
    }


def _infer_delimiter(path: Path) -> str:
    if path.suffix.lower() in {".tsv", ".txt"}:
        return "\t"
    try:
        sample = path.read_text(encoding="utf-8-sig", errors="ignore")[:4096]
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except (OSError, csv.Error):
        return ","
    return str(dialect.delimiter)


def _likely_metabolomics_columns(header: list[str]) -> list[str]:
    normalized = {name.strip().lower(): name for name in header}
    markers = {
        "metabolite_identification",
        "chemical_formula",
        "chemical_shift",
        "smiles",
        "inchi",
        "smallmolecule_abundance_sub",
    }
    return [normalized[item] for item in markers if item in normalized]


def _inspect_workbook(source_path: Path, metadata_path: Path | None) -> dict[str, Any]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12), values_only=True))
    raw_sample_metadata = _read_sample_metadata(metadata_path) if metadata_path and metadata_path.exists() else []
    sample_metadata = _dedupe_metadata_by_sample(raw_sample_metadata)
    detected_header_row = 0
    detected_header: list[str] = []
    duplicate_sample_columns: dict[str, list[int]] = {}
    numeric_column_ratio = 0.0
    try:
        detected_header_row, detected_header = _find_expression_header_row(sheet)
        annotation_start = _find_index(detected_header, "GeneID") or len(detected_header)
        preview = [
            [_cell_to_text(cell) for cell in row]
            for row in sheet.iter_rows(
                min_row=detected_header_row + 1,
                max_row=min(sheet.max_row, detected_header_row + 8),
                values_only=True,
            )
        ]
        duplicate_sample_columns = {
            name: indices
            for name, indices in _all_column_indices(detected_header[:annotation_start]).items()
            if len(indices) > 1 and name not in GENE_COLUMNS
        }
        ratios = _numeric_ratios(preview, annotation_start)
        numeric_column_ratio = _safe_ratio(sum(ratio > 0.8 for ratio in ratios), len(ratios))
    except ValueError:
        pass
    result = {
        "file_type": "workbook",
        "path": str(source_path),
        "sheet_names": workbook.sheetnames,
        "first_sheet": sheet.title,
        "row_count": sheet.max_row,
        "column_count": sheet.max_column,
        "preview": [[_cell_to_text(cell) for cell in row[:30]] for row in rows[:8]],
        "metadata_rows": len(sample_metadata),
        "metadata_conditions": _count_values(row.get("condition", "") for row in sample_metadata),
        "detected_header_row": detected_header_row,
        "header": detected_header[:160],
        "likely_gene_columns": [name for name in GENE_COLUMNS if name in detected_header],
        "duplicate_sample_columns": duplicate_sample_columns,
        "numeric_column_ratio": numeric_column_ratio,
    }
    workbook.close()
    return result


def _classify_with_deepseek(inspection: dict[str, Any], use_llm: bool) -> dict[str, Any]:
    fallback = _heuristic_plan(inspection)
    if not use_llm:
        return fallback | {"llm_status": "disabled"}

    api_key = _env_value("DEEPSEEK_API_KEY")
    if not api_key:
        return fallback | {"llm_status": "missing_api_key"}

    base_url = _env_value("DEEPSEEK_BASE_URL", "https://api.deepseek.com").rstrip("/")
    model = _env_value("DEEPSEEK_ROUTER_MODEL", DEFAULT_MODEL)
    user_prompt = "请返回 JSON。inspection JSON:\n" + json.dumps(
        _compact_for_prompt(inspection),
        ensure_ascii=False,
    )
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": DATA_INTAKE_SYSTEM_PROMPT},
            {"role": "user", "content": user_prompt},
        ],
        "response_format": {"type": "json_object"},
        "max_tokens": 1200,
        "temperature": 0.1,
    }
    request = urllib.request.Request(
        f"{base_url}/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=20) as response:
            raw = json.loads(response.read().decode("utf-8"))
        content = raw["choices"][0]["message"].get("content") or ""
        if not content.strip():
            return fallback | {"llm_status": "empty_content", "model": model}
        parsed = json.loads(content)
        return fallback | parsed | {"llm_status": "ok", "model": model}
    except (urllib.error.URLError, TimeoutError, KeyError, json.JSONDecodeError) as exc:
        return fallback | {"llm_status": "fallback", "llm_error": str(exc), "model": model}


def _standardize_expression_like_table(
    source_path: Path,
    metadata_path: Path | None,
    output_dir: Path,
    inspection: dict[str, Any],
    llm_plan: dict[str, Any],
    strategy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    if source_path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _standardize_workbook_expression_table(
            source_path=source_path,
            metadata_path=metadata_path,
            output_dir=output_dir,
            inspection=inspection,
            llm_plan=llm_plan,
            strategy=strategy,
        )

    matrix_path = output_dir / "expression_matrix.csv"
    sample_meta_path = output_dir / "sample_metadata.csv"
    annotation_path = output_dir / "gene_annotations.csv"
    raw_sample_metadata = _read_sample_metadata(metadata_path) if metadata_path and metadata_path.exists() else []
    sample_metadata = _dedupe_metadata_by_sample(raw_sample_metadata)

    with source_path.open(encoding="utf-8-sig", newline="") as source_file:
        reader = csv.reader(source_file)
        header = next(reader)
        preview = [row for _, row in zip(range(20), reader)]
        source_file.seek(0)
        reader = csv.reader(source_file)
        header = next(reader)
        sample_mode = str((strategy or {}).get("sample_mode") or "auto")
        duplicate_policy = str((strategy or {}).get("duplicate_policy") or "last")
        if sample_metadata and sample_mode == "metadata":
            sample_indices = _sample_column_indices(header, [row["sample"] for row in sample_metadata])
            selected_samples = []
            for row in sample_metadata:
                indices = sample_indices.get(row["sample"], [])
                if indices:
                    selected_index = indices[-1] if duplicate_policy == "last" else indices[0]
                    selected_samples.append((row["sample"], selected_index))
        else:
            selected_samples = _infer_numeric_sample_columns(header, preview, duplicate_policy=duplicate_policy)

        if not selected_samples:
            raise ValueError("No usable expression sample columns were detected")

        gene_indices = [idx for idx, name in enumerate(header) if name in GENE_COLUMNS]
        gene_name_to_index = {header[idx]: idx for idx in gene_indices}
        output_gene_columns = GENE_COLUMNS
        selected_samples = _dedupe_selected_sample_names(selected_samples)
        output_header = output_gene_columns + [name for name, _ in selected_samples]

        with (
            matrix_path.open("w", encoding="utf-8-sig", newline="") as matrix_file,
            annotation_path.open("w", encoding="utf-8-sig", newline="") as annotation_file,
        ):
            matrix_writer = csv.writer(matrix_file)
            annotation_writer = csv.writer(annotation_file)
            matrix_writer.writerow(output_header)
            annotation_writer.writerow(output_gene_columns)
            for row in reader:
                if not row or len(row) <= max(idx for _, idx in selected_samples):
                    continue
                gene_values = [
                    _value_at(row, gene_name_to_index[column])
                    if column in gene_name_to_index
                    else ""
                    for column in output_gene_columns
                ]
                sample_values = [_value_at(row, idx) for _, idx in selected_samples]
                if not any(gene_values):
                    continue
                matrix_writer.writerow(gene_values + sample_values)
                annotation_writer.writerow(gene_values)

    with sample_meta_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["sample", "group", "condition"])
        metadata_by_sample = {row["sample"]: row for row in sample_metadata}
        for sample, _ in selected_samples:
            row = metadata_by_sample.get(sample, {})
            inferred_group = _infer_group_from_sample_name(sample)
            group = row.get("group") or inferred_group
            condition = row.get("condition") or inferred_group
            writer.writerow([sample, group, condition])

    return {
        "matrix_file": str(matrix_path),
        "sample_metadata_file": str(sample_meta_path),
        "gene_annotation_file": str(annotation_path),
        "standardization": {
            "mode": "expression_matrix_from_metadata_samples",
            "sample_column_strategy": f"{sample_mode}_{duplicate_policy}_duplicate_column",
            "selected_sample_count": len(selected_samples),
            "gene_column_count": REQUIRED_MATRIX_GENE_COLUMNS,
            "llm_strategy": llm_plan.get("sample_columns_strategy", ""),
            "duplicate_sample_columns": inspection.get("duplicate_sample_columns", {}),
            "strategy_id": (strategy or {}).get("id", ""),
        },
    }


def _standardize_workbook_expression_table(
    source_path: Path,
    metadata_path: Path | None,
    output_dir: Path,
    inspection: dict[str, Any],
    llm_plan: dict[str, Any],
    strategy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    matrix_path = output_dir / "expression_matrix.csv"
    sample_meta_path = output_dir / "sample_metadata.csv"
    annotation_path = output_dir / "gene_annotations.csv"
    raw_sample_metadata = _read_sample_metadata(metadata_path) if metadata_path and metadata_path.exists() else []
    sample_metadata = _dedupe_metadata_by_sample(raw_sample_metadata)

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row_number, header = _find_expression_header_row(sheet)
    group_row = _workbook_row_text(sheet, header_row_number - 1) if header_row_number > 1 else []
    annotation_start = _find_index(header, "GeneID") or len(header)

    sample_mode = str((strategy or {}).get("sample_mode") or "auto")
    duplicate_policy = str((strategy or {}).get("duplicate_policy") or "last")
    if sample_metadata and sample_mode == "metadata":
        selected_samples = []
        sample_indices = _sample_column_indices(header[:annotation_start], [row["sample"] for row in sample_metadata])
        for row in sample_metadata:
            indices = sample_indices.get(row["sample"], [])
            if indices:
                selected_index = indices[-1] if duplicate_policy == "last" else indices[0]
                selected_samples.append((row["sample"], selected_index))
    else:
        preview = [
            [_cell_to_text(cell) for cell in row]
            for row in sheet.iter_rows(
                min_row=header_row_number + 1,
                max_row=min(sheet.max_row, header_row_number + 20),
                values_only=True,
            )
        ]
        selected_samples = _infer_numeric_sample_columns(
            header[:annotation_start],
            preview,
            duplicate_policy=duplicate_policy,
        )

    if not selected_samples:
        raise ValueError("No usable expression sample columns were detected in workbook")

    gene_indices = [idx for idx, name in enumerate(header) if name in GENE_COLUMNS]
    gene_name_to_index = {header[idx]: idx for idx in gene_indices}
    selected_samples = _dedupe_selected_sample_names(selected_samples)
    output_header = GENE_COLUMNS + [name for name, _ in selected_samples]
    annotation_columns = [
        (name, idx)
        for idx, name in enumerate(header[annotation_start:], start=annotation_start)
        if name
    ]

    with (
        matrix_path.open("w", encoding="utf-8-sig", newline="") as matrix_file,
        annotation_path.open("w", encoding="utf-8-sig", newline="") as annotation_file,
    ):
        matrix_writer = csv.writer(matrix_file)
        annotation_writer = csv.writer(annotation_file)
        matrix_writer.writerow(output_header)
        annotation_writer.writerow(GENE_COLUMNS + [name for name, _ in annotation_columns])

        for row in sheet.iter_rows(min_row=header_row_number + 1, values_only=True):
            if not row or len(row) <= max(idx for _, idx in selected_samples):
                continue
            gene_values = [
                _cell_to_text(_value_at(row, gene_name_to_index[column]))
                if column in gene_name_to_index
                else ""
                for column in GENE_COLUMNS
            ]
            if not any(gene_values):
                continue
            sample_values = [_cell_to_text(_value_at(row, idx)) for _, idx in selected_samples]
            annotation_values = [_cell_to_text(_value_at(row, idx)) for _, idx in annotation_columns]
            matrix_writer.writerow(gene_values + sample_values)
            annotation_writer.writerow(gene_values + annotation_values)

    workbook_conditions = _workbook_conditions_for_samples(group_row, selected_samples)
    with sample_meta_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["sample", "group", "condition"])
        metadata_by_sample = {row["sample"]: row for row in sample_metadata}
        for sample, _ in selected_samples:
            row = metadata_by_sample.get(sample, {})
            inferred_group = _infer_group_from_sample_name(sample)
            inferred_condition = workbook_conditions.get(sample) or inferred_group
            group = row.get("group") or inferred_group
            condition = row.get("condition") or inferred_condition
            writer.writerow([sample, group, condition])

    workbook.close()
    return {
        "matrix_file": str(matrix_path),
        "sample_metadata_file": str(sample_meta_path),
        "gene_annotation_file": str(annotation_path),
        "standardization": {
            "mode": "expression_matrix_from_workbook",
            "sheet_name": sheet.title,
            "header_row": header_row_number,
            "sample_column_strategy": f"{sample_mode}_{duplicate_policy}_duplicate_numeric_column_before_annotation",
            "selected_sample_count": len(selected_samples),
            "gene_column_count": REQUIRED_MATRIX_GENE_COLUMNS,
            "llm_strategy": llm_plan.get("sample_columns_strategy", ""),
            "duplicate_sample_columns": inspection.get("duplicate_sample_columns", {}),
            "strategy_id": (strategy or {}).get("id", ""),
        },
    }


def _standardize_metabolomics_table(
    source_path: Path,
    metadata_path: Path | None,
    output_dir: Path,
    inspection: dict[str, Any],
    llm_plan: dict[str, Any],
    strategy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    matrix_path = output_dir / "metabolomics_matrix.csv"
    sample_meta_path = output_dir / "sample_metadata.csv"
    annotation_path = output_dir / "metabolite_annotations.csv"
    delimiter = str(inspection.get("delimiter") or _infer_delimiter(source_path))
    metadata_rows = _read_metabolomics_metadata(metadata_path) or _read_metabolomics_metadata(
        _find_sibling_metabolights_sample_file(source_path)
    )

    with source_path.open(encoding="utf-8-sig", newline="") as source_file:
        reader = csv.reader(source_file, delimiter=delimiter)
        header = next(reader, [])
        preview = [row for _, row in zip(range(24), reader)]
        source_file.seek(0)
        reader = csv.reader(source_file, delimiter=delimiter)
        next(reader, None)

        sample_names = [row["sample"] for row in metadata_rows]
        selected_samples = _metabolomics_sample_columns(header, preview, sample_names)
        if not selected_samples:
            raise ValueError("No usable metabolomics sample columns were detected")

        header_lookup = {name.strip().lower(): index for index, name in enumerate(header)}
        annotation_fields = [
            "database_identifier",
            "chemical_formula",
            "smiles",
            "inchi",
            "metabolite_identification",
            "chemical_shift",
            "taxid",
            "species",
            "database",
            "reliability",
            "uri",
        ]

        with (
            matrix_path.open("w", encoding="utf-8-sig", newline="") as matrix_file,
            annotation_path.open("w", encoding="utf-8-sig", newline="") as annotation_file,
        ):
            matrix_writer = csv.writer(matrix_file)
            annotation_writer = csv.writer(annotation_file)
            matrix_writer.writerow(GENE_COLUMNS + [name for name, _ in selected_samples])
            annotation_writer.writerow(["feature_id", "feature_name", *annotation_fields])

            metabolite_count = 0
            for index, row in enumerate(reader, start=1):
                if not row:
                    continue
                values = [_value_at(row, column_index) for _, column_index in selected_samples]
                if not any(_is_number(value) for value in values):
                    continue
                feature_name = _first_text(
                    _value_by_header(row, header_lookup, "metabolite_identification"),
                    _value_by_header(row, header_lookup, "chemical_shift"),
                    f"metabolite_{index}",
                )
                feature_id = _first_text(
                    _value_by_header(row, header_lookup, "database_identifier"),
                    feature_name,
                    f"metabolite_{index}",
                )
                formula = _value_by_header(row, header_lookup, "chemical_formula")
                locus = _first_text(
                    _value_by_header(row, header_lookup, "uri"),
                    _value_by_header(row, header_lookup, "smiles"),
                    formula,
                )
                matrix_writer.writerow([feature_name, feature_id, "metabolite", "", locus, "1", *values])
                annotation_writer.writerow(
                    [
                        feature_id,
                        feature_name,
                        *[_value_by_header(row, header_lookup, field) for field in annotation_fields],
                    ]
                )
                metabolite_count += 1

    metadata_by_sample = {row["sample"]: row for row in metadata_rows}
    with sample_meta_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["sample", "group", "condition"])
        for sample, _ in selected_samples:
            row = metadata_by_sample.get(sample, {})
            condition = row.get("condition") or _infer_group_from_sample_name(sample)
            group = row.get("group") or condition
            writer.writerow([sample, group, condition])

    return {
        "matrix_file": str(matrix_path),
        "sample_metadata_file": str(sample_meta_path),
        "gene_annotation_file": str(annotation_path),
        "standardization": {
            "mode": "metabolomics_matrix_from_maf",
            "selected_sample_count": len(selected_samples),
            "metabolite_count": metabolite_count,
            "metadata_rows": len(metadata_rows),
            "llm_strategy": llm_plan.get("sample_columns_strategy", ""),
            "strategy_id": (strategy or {}).get("id", ""),
        },
    }


def _metabolomics_sample_columns(
    header: list[str],
    preview: list[list[str]],
    sample_names: list[str],
) -> list[tuple[str, int]]:
    if sample_names:
        sample_indices = _sample_column_indices(header, sample_names)
        selected = [(sample, indices[0]) for sample, indices in sample_indices.items() if indices]
        if selected:
            return selected

    annotation_names = {
        "database_identifier",
        "chemical_formula",
        "smiles",
        "inchi",
        "metabolite_identification",
        "chemical_shift",
        "multiplicity",
        "taxid",
        "species",
        "database",
        "database_version",
        "reliability",
        "uri",
        "search_engine",
        "search_engine_score",
        "smallmolecule_abundance_sub",
        "smallmolecule_abundance_stdev_sub",
        "smallmolecule_abundance_std_error_sub",
    }
    selected = []
    for index, name in enumerate(header):
        if name.strip().lower() in annotation_names:
            continue
        values = [_value_at(row, index) for row in preview if index < len(row)]
        if values and _safe_ratio(sum(_is_number(value) for value in values), len(values)) >= 0.7:
            selected.append((name or f"sample_{index + 1}", index))
    return selected


def _read_metabolomics_metadata(path: Path | None) -> list[dict[str, str]]:
    rows = _read_sample_metadata(path)
    if rows:
        return rows
    if path is None or not path.exists():
        return []
    delimiter = _infer_delimiter(path)
    with path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, delimiter=delimiter)
        if not reader.fieldnames:
            return []
        field_lookup = {_metadata_key(name): name for name in reader.fieldnames}
        sample_key = field_lookup.get("sample_name") or field_lookup.get("sample")
        condition_key = (
            field_lookup.get("factor_value_metabolic_syndrome")
            or field_lookup.get("factor_value_condition")
            or field_lookup.get("condition")
            or field_lookup.get("characteristics_sample_type")
        )
        group_key = field_lookup.get("factor_value_gender") or field_lookup.get("source_name")
        if sample_key is None:
            return []
        parsed = []
        for raw in reader:
            sample = str(raw.get(sample_key) or "").strip()
            if not sample:
                continue
            condition = str(raw.get(condition_key) or "unknown").strip() if condition_key else "unknown"
            group = str(raw.get(group_key) or condition).strip() if group_key else condition
            parsed.append({"sample": sample, "group": group or condition, "condition": condition or "unknown"})
        return parsed


def _find_sibling_metabolights_sample_file(source_path: Path) -> Path | None:
    candidates = sorted(source_path.parent.glob("s_*.txt"))
    return candidates[0] if candidates else None


def _metadata_key(value: str | None) -> str:
    return "".join(char.lower() if char.isalnum() else "_" for char in str(value or "").strip()).strip("_")


def _value_by_header(row: list[str], header_lookup: dict[str, int], name: str) -> str:
    index = header_lookup.get(name)
    if index is None:
        return ""
    return str(_value_at(row, index)).strip()


def _first_text(*values: str) -> str:
    for value in values:
        text = str(value or "").strip()
        if text:
            return text
    return ""


def _validate_expression_matrix(matrix_path: Path, metadata_path: Path) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    if not matrix_path.exists():
        errors.append("standard matrix file missing")
    if not metadata_path.exists():
        errors.append("sample metadata file missing")
    if errors:
        return {"valid": False, "errors": errors, "warnings": warnings}

    with matrix_path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        header = next(reader, [])
        preview = [row for _, row in zip(range(20), reader)]
    metadata = _read_sample_metadata(metadata_path)
    sample_names = [row["sample"] for row in metadata]
    sample_indices = [header.index(sample) for sample in sample_names if sample in header]
    if not sample_indices:
        errors.append("no metadata samples matched matrix columns")
    if len(sample_indices) < 2:
        errors.append("expression matrix requires at least 2 sample columns")

    numeric_ok = 0
    numeric_total = 0
    for row in preview:
        for idx in sample_indices:
            if idx >= len(row):
                continue
            numeric_total += 1
            if _is_number(row[idx]):
                numeric_ok += 1
    numeric_ratio = _safe_ratio(numeric_ok, numeric_total)
    if numeric_ratio < 0.8:
        errors.append("sample columns are not numeric enough for expression analysis")

    conditions = _count_values(row["condition"] for row in metadata)
    groups = _count_values(row["group"] for row in metadata)
    if len(conditions) < 2:
        warnings.append("metadata has fewer than 2 conditions; diff analysis will be unavailable")

    gene_count = _count_csv_rows(matrix_path)
    return {
        "valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "gene_count": gene_count,
        "sample_count": len(sample_indices),
        "sample_groups": groups,
        "conditions": conditions,
        "numeric_sample_value_ratio": numeric_ratio,
    }


def _capabilities_for_validation(validation: dict[str, Any]) -> list[str]:
    capabilities = []
    if validation["valid"] and validation["sample_count"] >= 2:
        capabilities.extend(["qc", "sample_correlation", "expression_heatmap", "gene_expression"])
        capabilities.append("pca")
    if validation["valid"] and validation["sample_count"] > 20:
        capabilities.append("wgcna")
    condition_counts = validation.get("conditions", {})
    if len(condition_counts) >= 2 and all(count >= 2 for count in condition_counts.values()):
        capabilities.append("diff_analysis")
    if len(condition_counts) >= 3:
        capabilities.append("multigroup_differential")
    return capabilities


def _next_analyses_for_capabilities(capabilities: list[str]) -> list[dict[str, str]]:
    if "qc" in capabilities:
        return [
            {
                "type": "qc",
                "label": "Multi-sample QC",
                "description": "Gate expression data through sample QC before downstream analysis.",
            }
        ]
    specs = {
        "qc": {
            "type": "qc",
            "label": "矩阵 QC",
            "description": "查看表达量分布、总量和零值比例。",
        },
        "sample_correlation": {
            "type": "sample_correlation",
            "label": "样本相关性",
            "description": "查看样本间相关性热图。",
        },
        "expression_heatmap": {
            "type": "expression_heatmap",
            "label": "表达热图",
            "description": "基于高变基因生成表达聚类热图。",
        },
        "gene_expression": {
            "type": "gene_expression",
            "label": "单基因表达",
            "description": "查看指定基因在不同分组中的表达。",
        },
        "pca": {
            "type": "pca",
            "label": "PCA",
            "description": "基于表达矩阵查看样本整体分布。",
        },
        "diff_analysis": {
            "type": "diff_analysis",
            "label": "差异分析",
            "description": "选择两个样本分组进行差异表达分析。",
        },
        "metabolomics_normalization": {
            "type": "metabolomics_normalization",
            "label": "Normalize / impute / scale",
            "description": "Prepare metabolomics intensities for downstream exploratory and differential analysis.",
        },
        "metabolomics_differential": {
            "type": "metabolomics_differential",
            "label": "Metabolomics differential analysis",
            "description": "Compare two metabolomics conditions and report differential metabolites.",
        },
    }
    return [specs[item] for item in capabilities if item in specs]


def _heuristic_plan(inspection: dict[str, Any]) -> dict[str, Any]:
    likely_metabolomics = bool(inspection.get("likely_metabolomics_columns"))
    likely_expression = bool(inspection.get("likely_gene_columns")) and (
        inspection.get("metadata_rows", 0) > 0 or inspection.get("numeric_column_ratio", 0) > 0.5
    )
    likely_single_cell = bool(inspection.get("likely_gene_columns")) and inspection.get("column_count", 0) > 1000
    likely_feature_table = not inspection.get("likely_gene_columns") and inspection.get("numeric_column_ratio", 0) > 0.6
    if likely_metabolomics:
        data_type = "metabolomics_matrix"
        capabilities = ["pca", "metabolomics_normalization", "metabolomics_differential"]
    elif likely_single_cell:
        data_type = "single_cell_matrix"
        capabilities: list[str] = []
    elif likely_expression:
        data_type = "expression_matrix"
        capabilities = ["pca", "diff_analysis"]
    elif likely_feature_table:
        data_type = "feature_table"
        capabilities = []
    else:
        data_type = "unknown_table"
        capabilities = []
    return {
        "data_type": data_type,
        "confidence": 0.85 if data_type != "unknown_table" else 0.4,
        "feature_id_column": "gene_id",
        "feature_name_column": "gene_short_name",
        "sample_columns_strategy": "use_metadata_samples_last_duplicate",
        "needs_metadata": True,
        "capabilities": capabilities,
        "warnings": [],
        "model": DEFAULT_MODEL,
    }


def _record_step(
    checkpoint_path: Path,
    checkpoint: dict[str, Any],
    name: str,
    status: str,
    payload: dict[str, Any],
) -> None:
    checkpoint["iterations"].append(
        {
            "step": name,
            "status": status,
            "payload": payload,
            "timestamp": datetime.now().isoformat(),
        }
    )
    checkpoint["updated_at"] = datetime.now().isoformat()
    _write_checkpoint(checkpoint_path, checkpoint)


def _record_attempt(
    checkpoint_path: Path,
    checkpoint: dict[str, Any],
    payload: dict[str, Any],
) -> None:
    checkpoint["attempts"].append(payload)
    checkpoint["updated_at"] = datetime.now().isoformat()
    _write_checkpoint(checkpoint_path, checkpoint)


def _write_checkpoint(path: Path, checkpoint: dict[str, Any]) -> None:
    path.write_text(json.dumps(checkpoint, ensure_ascii=False, indent=2), encoding="utf-8")


def _emit_progress(
    callback: Callable[[str, str, str], None] | None,
    step: str,
    status: str,
    label: str,
) -> None:
    if callback is None:
        return
    callback(step, status, label)


def _failure_summary(inspection: dict[str, Any], validation: dict[str, Any]) -> str:
    file_type = inspection.get("file_type") or "unknown"
    column_count = inspection.get("column_count", 0)
    row_count = inspection.get("row_count", 0)
    if validation.get("errors"):
        return "文件已被读取，但标准化后的样本列或分组信息不满足当前分析流程要求。"
    if inspection.get("likely_gene_columns"):
        return (
            f"文件已读取为 {file_type}，共 {row_count} 行、{column_count} 列，"
            "但没有成功定位出可直接用于表达分析的样本表达列。"
        )
    return (
        f"文件已读取为 {file_type}，共 {row_count} 行、{column_count} 列，"
        "但整体结构不像当前支持的 bulk RNA 表达矩阵。"
    )


def _failure_findings(inspection: dict[str, Any], validation: dict[str, Any]) -> list[str]:
    findings: list[str] = []
    gene_columns = inspection.get("likely_gene_columns", [])
    if gene_columns:
        findings.append(f"识别到的基因注释列：{', '.join(gene_columns)}")
    else:
        findings.append("没有识别到表达矩阵常见的基因注释列，如 gene_short_name、gene_id、Length。")

    numeric_ratio = float(inspection.get("numeric_column_ratio", 0) or 0)
    if numeric_ratio:
        findings.append(f"整表数值列占比约为 {numeric_ratio:.0%}。")

    sample_numeric_ratio = float(inspection.get("numeric_sample_column_ratio", 0) or 0)
    if sample_numeric_ratio:
        findings.append(f"候选样本列中的数值占比约为 {sample_numeric_ratio:.0%}。")

    duplicate_samples = inspection.get("duplicate_sample_columns", {})
    if duplicate_samples:
        findings.append(f"发现重复样本列 {len(duplicate_samples)} 组，Agent 会优先尝试使用后出现的一组。")

    for error in validation.get("errors", []):
        findings.append(f"验证结果：{error}")
    return findings


def _failure_reasons(
    error_message: str,
    inspection: dict[str, Any],
    validation: dict[str, Any],
) -> list[str]:
    reasons: list[str] = []
    if "No usable expression sample columns were detected" in error_message:
        reasons.append("Agent 没有在基因注释列之后找到足够稳定的数值样本列。")
        reasons.append("这通常意味着样本表达值不是按列摆放，或者列里混入了大量文字/备注，无法当作表达矩阵。")
    if "Could not locate workbook expression header row" in error_message:
        reasons.append("工作表中没有找到包含 gene_short_name、gene_id、Length 等关键列名的表头。")
    if "expression matrix requires at least 2 sample columns" in error_message:
        reasons.append("当前可识别的表达样本列少于 2 列，无法进行 PCA 或差异分析。")
    if "sample columns are not numeric enough" in error_message:
        reasons.append("识别到的候选样本列里，数值比例过低，说明这些列不是纯表达量列。")
    if "no metadata samples matched matrix columns" in error_message:
        reasons.append("样本分组表中的 sample 名称和表达矩阵列名对不上。")

    if not reasons:
        if not inspection.get("likely_gene_columns"):
            reasons.append("文件结构不像当前支持的 bulk RNA 表达矩阵。")
        elif validation.get("errors"):
            reasons.append("标准化后仍未通过当前流程的表达矩阵校验。")
        else:
            reasons.append(error_message)
    return reasons


def _failure_suggestions(
    error_message: str,
    inspection: dict[str, Any],
    validation: dict[str, Any],
) -> list[str]:
    suggestions: list[str] = []
    if not inspection.get("likely_gene_columns"):
        suggestions.append("确认文件里有一行真实表头，并包含 gene_short_name、gene_id、Length 这类基因注释列。")
    if "No usable expression sample columns were detected" in error_message:
        suggestions.append("确认样本是按列排列，每个样本列都应主要由数值表达量组成。")
        suggestions.append("去掉说明文字、合并单元格影响区、备注列和统计汇总列后再上传。")
    if "sample columns are not numeric enough" in error_message:
        suggestions.append("只保留表达量矩阵本身，不要混入注释文本列。")
    if "no metadata samples matched matrix columns" in error_message:
        suggestions.append("如果你后续单独提供分组信息，sample 列必须和表达矩阵列名完全一致。")
    if "expression matrix requires at least 2 sample columns" in error_message:
        suggestions.append("至少需要 2 个样本列，且更适合每组不少于 2 个样本。")
    if not suggestions and validation.get("errors"):
        suggestions.append("优先检查样本列是否为纯数值、分组列名是否一致。")
    if not suggestions:
        suggestions.append("如果这是单细胞矩阵或其他非 bulk RNA 表格，需要补对应的数据标准化器后再支持。")
    return suggestions




def _env_value(name: str, default: str = "") -> str:
    if name in os.environ:
        return os.environ[name]
    env_path = PROJECT_ROOT / ".env"
    if not env_path.exists():
        return default
    for line in env_path.read_text(encoding="utf-8").splitlines():
        if not line or line.lstrip().startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        if key.strip() == name:
            return value.strip().strip('"').strip("'")
    return default


def _compact_for_prompt(inspection: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in inspection.items()
        if key not in {"preview"} or isinstance(value, (str, int, float))
    } | {"preview": inspection.get("preview", [])[:2]}


def _read_sample_metadata(path: Path | None) -> list[dict[str, str]]:
    if path is None or not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    required = {"sample", "group", "condition"}
    if not rows or not required.issubset(rows[0]):
        return []
    return rows


def _dedupe_metadata_by_sample(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    by_sample: dict[str, dict[str, str]] = {}
    for row in rows:
        sample = row.get("sample", "")
        if sample:
            by_sample[sample] = row
    return list(by_sample.values())


def _infer_numeric_sample_columns(
    header: list[str],
    preview: list[list[str]],
    duplicate_policy: str = "last",
) -> list[tuple[str, int]]:
    gene_indices = {index for index, name in enumerate(header) if name in GENE_COLUMNS}
    inferred: dict[str, int] = {}
    for index, name in enumerate(header):
        if index in gene_indices:
            continue
        values = [_value_at(row, index) for row in preview if index < len(row)]
        if values and _safe_ratio(sum(_is_number(value) for value in values), len(values)) >= 0.8:
            key = name or f"sample_{index + 1}"
            if duplicate_policy == "first" and key in inferred:
                continue
            inferred[key] = index
    return list(inferred.items())


def _infer_group_from_sample_name(sample: str) -> str:
    name = sample.rsplit("_", 1)[0] if sample.rsplit("_", 1)[-1].isdigit() else sample
    if "-" in name:
        prefix = name.split("-", 1)[0].strip()
    elif "_" in name:
        prefix = name.split("_", 1)[0].strip()
    else:
        prefix = "unknown"
    prefix = prefix.removeprefix("Group ").strip()
    return prefix or "unknown"


def _find_expression_header_row(sheet: Any) -> tuple[int, list[str]]:
    max_scan_row = min(sheet.max_row, 80)
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_scan_row, values_only=True),
        start=1,
    ):
        header = [_cell_to_text(cell).strip() for cell in row]
        if all(column in header for column in GENE_COLUMNS):
            return row_number, header
    raise ValueError("Could not locate workbook expression header row")


def _workbook_row_text(sheet: Any, row_number: int) -> list[str]:
    row = next(sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True), [])
    return [_cell_to_text(cell).strip() for cell in row]


def _workbook_conditions_for_samples(
    group_row: list[str],
    selected_samples: list[tuple[str, int]],
) -> dict[str, str]:
    current = ""
    index_to_condition: dict[int, str] = {}
    max_index = max((index for _, index in selected_samples), default=-1)
    for index in range(max_index + 1):
        value = group_row[index] if index < len(group_row) else ""
        if value:
            current = _normalize_condition_label(value)
        if current:
            index_to_condition[index] = current
    return {sample: index_to_condition.get(index, "") for sample, index in selected_samples}


def _normalize_condition_label(value: str) -> str:
    normalized = value.strip().lower()
    aliases = {
        "nor": "normal",
        "normal": "normal",
        "ca": "cancer",
        "cancer": "cancer",
        "tumor": "cancer",
        "lm": "lm",
    }
    return aliases.get(normalized, normalized or "unknown")


def _find_index(values: list[str], target: str) -> int | None:
    for index, value in enumerate(values):
        if value == target:
            return index
    return None


def _dedupe_selected_sample_names(samples: list[tuple[str, int]]) -> list[tuple[str, int]]:
    counts: dict[str, int] = {}
    deduped = []
    for sample, index in samples:
        counts[sample] = counts.get(sample, 0) + 1
        name = sample if counts[sample] == 1 else f"{sample}_{counts[sample]}"
        deduped.append((name, index))
    return deduped


def _sample_column_indices(header: list[str], sample_names: list[str]) -> dict[str, list[int]]:
    wanted = set(sample_names)
    mapping = {sample: [] for sample in sample_names}
    for index, name in enumerate(header):
        if name in wanted:
            mapping[name].append(index)
    return mapping


def _all_column_indices(header: list[str]) -> dict[str, list[int]]:
    mapping: dict[str, list[int]] = {}
    for index, name in enumerate(header):
        if not name:
            continue
        mapping.setdefault(name, []).append(index)
    return mapping


def _numeric_ratios(rows: list[list[str]], column_count: int) -> list[float]:
    ratios = []
    for idx in range(column_count):
        values = [_value_at(row, idx) for row in rows if idx < len(row)]
        ratios.append(_safe_ratio(sum(_is_number(value) for value in values), len(values)))
    return ratios


def _numeric_sample_ratio(rows: list[list[str]], sample_indices: dict[str, list[int]]) -> float:
    values = []
    for indices in sample_indices.values():
        values.extend(indices)
    total = 0
    numeric = 0
    for row in rows:
        for idx in values:
            if idx >= len(row):
                continue
            total += 1
            numeric += int(_is_number(row[idx]))
    return _safe_ratio(numeric, total)


def _duplicates(values: Iterable[str]) -> dict[str, int]:
    counts = _count_values(values)
    return {key: value for key, value in counts.items() if value > 1}


def _count_values(values: Iterable[str]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for value in values:
        key = str(value)
        counts[key] = counts.get(key, 0) + 1
    return counts


def _count_csv_rows(path: Path) -> int:
    with path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        next(reader, None)
        return sum(1 for row in reader if row)


def _value_at(row: list[str], index: int) -> str:
    return row[index] if index < len(row) else ""


def _is_number(value: Any) -> bool:
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def _safe_ratio(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def _cell_to_text(value: Any) -> str:
    return "" if value is None else str(value)
