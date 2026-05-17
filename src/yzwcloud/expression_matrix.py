from __future__ import annotations

import csv
import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from yzwcloud.config import PROJECT_ROOT
from yzwcloud.data_intake_agent import run_data_intake_agent
from yzwcloud.models import DataObject


DEFAULT_EXPRESSION_FILE = "expression_matrix.csv"
DEFAULT_SAMPLE_METADATA_FILE = "sample_metadata.csv"
DEFAULT_SHEET = "mRNA profiling"
GENE_INFO_COLUMNS = {
    "gene_short_name",
    "gene_id",
    "biotype",
    "strand",
    "locus",
    "Length",
}
ANNOTATION_START = "GeneID"


def prepare_expression_matrix(params: dict[str, Any], output_dir: Path) -> DataObject:
    source_path = _resolve_source_path(str(params.get("source_path") or DEFAULT_EXPRESSION_FILE))
    metadata_path = _optional_source_path(str(params.get("sample_metadata_path") or DEFAULT_SAMPLE_METADATA_FILE))
    if params.get("agent_enabled", True):
        return run_data_intake_agent(
            source_path=source_path,
            metadata_path=metadata_path,
            output_dir=output_dir,
            params=params,
        )

    if source_path.suffix.lower() == ".csv":
        return _prepare_csv_expression_matrix(params=params, source_path=source_path, output_dir=output_dir)

    sheet_name = str(params.get("sheet_name") or DEFAULT_SHEET)
    header_row = int(params.get("header_row") or 11)

    output_dir.mkdir(parents=True, exist_ok=True)
    matrix_path = output_dir / "expression_matrix.csv"
    sample_meta_path = output_dir / "sample_metadata.csv"
    annotation_path = output_dir / "gene_annotations.csv"

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")

    worksheet = workbook[sheet_name]
    group_row = _read_row(worksheet, header_row - 1)
    header = _read_row(worksheet, header_row)
    column_plan = _build_column_plan(header, group_row)

    gene_count = _write_normalized_files(
        worksheet=worksheet,
        header_row=header_row,
        column_plan=column_plan,
        matrix_path=matrix_path,
        sample_meta_path=sample_meta_path,
        annotation_path=annotation_path,
    )

    sample_count = len(column_plan["sample_columns"])
    metadata = {
        "source_file": str(source_path),
        "sheet_name": sheet_name,
        "header_row": header_row,
        "source_rows": worksheet.max_row,
        "source_columns": worksheet.max_column,
        "gene_count": gene_count,
        "sample_count": sample_count,
        "declared_sample_count": 208,
        "actual_sample_count_note": "文件说明写 208 samples；当前 sheet 表头实际解析出 104 个表达量样本列。",
        "sample_groups": _count_values(item["group"] for item in column_plan["sample_columns"]),
        "conditions": _count_values(item["condition"] for item in column_plan["sample_columns"]),
        "matrix_file": str(matrix_path),
        "sample_metadata_file": str(sample_meta_path),
        "gene_annotation_file": str(annotation_path),
    }

    return DataObject(type="expression_matrix", data=str(matrix_path), meta=metadata)


def _prepare_csv_expression_matrix(
    params: dict[str, Any],
    source_path: Path,
    output_dir: Path,
) -> DataObject:
    output_dir.mkdir(parents=True, exist_ok=True)
    matrix_path = output_dir / "expression_matrix.csv"
    sample_meta_path = output_dir / "sample_metadata.csv"

    shutil.copyfile(source_path, matrix_path)

    metadata_source = _resolve_source_path(
        str(params.get("sample_metadata_path") or DEFAULT_SAMPLE_METADATA_FILE)
    )
    shutil.copyfile(metadata_source, sample_meta_path)

    sample_rows = _read_sample_metadata(sample_meta_path)
    sample_count = len(sample_rows)
    gene_count = _count_csv_data_rows(matrix_path)

    metadata = {
        "source_file": str(source_path),
        "sample_metadata_source": str(metadata_source),
        "gene_count": gene_count,
        "sample_count": sample_count,
        "sample_groups": _count_values(row["group"] for row in sample_rows),
        "conditions": _count_values(row["condition"] for row in sample_rows),
        "condition_options": sorted({row["condition"] for row in sample_rows}),
        "matrix_file": str(matrix_path),
        "sample_metadata_file": str(sample_meta_path),
        "gene_annotation_file": "",
        "note": "使用根目录 CSV 表达矩阵和 sample_metadata.csv 作为演示输入。",
    }

    return DataObject(type="expression_matrix", data=str(matrix_path), meta=metadata)


def _resolve_source_path(value: str) -> Path:
    raw_path = Path(value)
    path = raw_path if raw_path.is_absolute() else PROJECT_ROOT / raw_path
    if not path.exists():
        raise FileNotFoundError(f"Expression matrix file not found: {path}")
    return path.resolve()


def _optional_source_path(value: str) -> Path | None:
    raw_path = Path(value)
    path = raw_path if raw_path.is_absolute() else PROJECT_ROOT / raw_path
    return path.resolve() if path.exists() else None


def _read_row(worksheet, row_number: int) -> list[Any]:
    return [cell.value for cell in next(worksheet.iter_rows(min_row=row_number, max_row=row_number))]


def _build_column_plan(header: list[Any], group_row: list[Any]) -> dict[str, list[dict[str, Any]]]:
    headers = [str(value).strip() if value is not None else "" for value in header]
    annotation_start_idx = _find_index(headers, ANNOTATION_START)
    if annotation_start_idx is None:
        raise ValueError(f"Cannot find annotation start column: {ANNOTATION_START}")

    gene_columns = [
        {"index": idx, "name": name}
        for idx, name in enumerate(headers[:annotation_start_idx])
        if name in GENE_INFO_COLUMNS
    ]
    sample_columns = []
    annotation_columns = [
        {"index": idx, "name": name}
        for idx, name in enumerate(headers[annotation_start_idx:], start=annotation_start_idx)
        if name
    ]

    current_condition = ""
    for idx in range(0, annotation_start_idx):
        header_name = headers[idx]
        condition_label = group_row[idx] if idx < len(group_row) else None
        if condition_label:
            current_condition = _normalize_condition(str(condition_label))
        if idx < len(gene_columns):
            continue
        if not header_name or not header_name.startswith("Group "):
            continue
        sample_columns.append(
            {
                "index": idx,
                "name": header_name,
                "group": _extract_group(header_name),
                "condition": current_condition or "unknown",
            }
        )

    if not sample_columns:
        raise ValueError("No expression sample columns found")

    return {
        "gene_columns": gene_columns,
        "sample_columns": sample_columns,
        "annotation_columns": annotation_columns,
    }


def _write_normalized_files(
    worksheet,
    header_row: int,
    column_plan: dict[str, list[dict[str, Any]]],
    matrix_path: Path,
    sample_meta_path: Path,
    annotation_path: Path,
) -> int:
    matrix_columns = column_plan["gene_columns"] + column_plan["sample_columns"]
    annotation_columns = column_plan["gene_columns"] + column_plan["annotation_columns"]

    with (
        matrix_path.open("w", encoding="utf-8-sig", newline="") as matrix_file,
        sample_meta_path.open("w", encoding="utf-8-sig", newline="") as sample_meta_file,
        annotation_path.open("w", encoding="utf-8-sig", newline="") as annotation_file,
    ):
        matrix_writer = csv.writer(matrix_file)
        sample_meta_writer = csv.writer(sample_meta_file)
        annotation_writer = csv.writer(annotation_file)

        matrix_writer.writerow([item["name"] for item in matrix_columns])
        annotation_writer.writerow([item["name"] for item in annotation_columns])
        sample_meta_writer.writerow(["sample", "group", "condition"])
        for sample in column_plan["sample_columns"]:
            sample_meta_writer.writerow([sample["name"], sample["group"], sample["condition"]])

        gene_count = 0
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            gene_id = _get_value(row, _index_by_name(column_plan["gene_columns"], "gene_id"))
            gene_name = _get_value(row, _index_by_name(column_plan["gene_columns"], "gene_short_name"))
            if gene_id is None and gene_name is None:
                continue

            matrix_writer.writerow([_get_value(row, item["index"]) for item in matrix_columns])
            annotation_writer.writerow([_get_value(row, item["index"]) for item in annotation_columns])
            gene_count += 1

    return gene_count


def _find_index(values: list[str], target: str) -> int | None:
    for idx, value in enumerate(values):
        if value == target:
            return idx
    return None


def _index_by_name(columns: list[dict[str, Any]], name: str) -> int | None:
    for column in columns:
        if column["name"] == name:
            return int(column["index"])
    return None


def _get_value(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _extract_group(sample_name: str) -> str:
    prefix = sample_name.split("-", 1)[0]
    return prefix.replace("Group ", "").strip() or "unknown"


def _normalize_condition(value: str) -> str:
    normalized = value.strip().lower()
    if normalized == "nor":
        return "normal"
    if normalized == "ca":
        return "cancer"
    return normalized


def _count_values(values) -> dict[str, int]:
    counts: dict[str, int] = {}
    for value in values:
        key = str(value)
        counts[key] = counts.get(key, 0) + 1
    return counts


def _read_sample_metadata(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    required = {"sample", "group", "condition"}
    if not rows or not required.issubset(rows[0]):
        raise ValueError("sample_metadata.csv must contain sample, group, condition columns")
    return rows


def _count_csv_data_rows(path: Path) -> int:
    with path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        next(reader, None)
        return sum(1 for row in reader if row)
